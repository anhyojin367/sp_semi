from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import linear_kernel

from .config import RAG_DATA_DIR, UCUM_JSON_CANDIDATES, UCUM_JSONL_CANDIDATES, UCUM_XLSX_CANDIDATES
from .utils import clean_text

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None


@dataclass
class RagDoc:
    idx: int
    text: str
    meta: dict


class UcumRagStore:
    def __init__(self) -> None:
        self.docs: list[RagDoc] = []
        self.word_vectorizer: TfidfVectorizer | None = None
        self.char_vectorizer: TfidfVectorizer | None = None
        self.word_matrix = None
        self.char_matrix = None
        self.loaded_sources: list[str] = []
        self._load()

    def _existing_paths(self, candidates: list[Path]) -> list[Path]:
        out: list[Path] = []
        seen: set[Path] = set()
        for path in candidates:
            if not path.exists():
                continue
            resolved = path.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            out.append(path)
        return out

    def _normalize_for_search(self, text: str | None) -> str:
        text = clean_text(text)
        if not text:
            return ""

        replacements = {
            "µ": "u",
            "μ": "u",
            "Å": "Ao",
            "Å": "Ao",
            "℃": "°C",
            "㎍": "ug",
            "㎎": "mg",
            "㎖": "mL",
            "㎕": "uL",
            "㎛": "um",
            "·": ".",
            "⋅": ".",
            "∙": ".",
            "⁻": "-",
            "−": "-",
            "²": "2",
            "³": "3",
            "¹": "1",
        }
        for src, dst in replacements.items():
            text = text.replace(src, dst)

        text = text.replace("mmHg", "mm[Hg] mmHg")
        text = text.replace("µL", "uL µL")
        text = text.replace("µg", "ug µg")
        text = text.replace("µm", "um µm")
        text = text.replace("µmol", "umol µmol")
        text = text.replace("µS/cm", "uS/cm µS/cm")
        return clean_text(text)

    def _append_loaded_source(self, source_name: str) -> None:
        if source_name and source_name not in self.loaded_sources:
            self.loaded_sources.append(source_name)

    def _rag_domain_for_path(self, path: Path) -> str:
        name = clean_text(path.stem).lower()

        if any(token in name for token in ("허가", "품목허가", "permit", "approval", "license")):
            return "허가서"

        if any(token in name for token in ("기호", "단위", "ucum", "symbol", "unit")):
            return "기호 사전"

        if any(
            token in name
            for token in (
                "생물학적",
                "기준 및 시험방법",
                "기준및시험방법",
                "별표",
                "약전",
                "통칙",
                "제제총칙",
                "의약품각조",
                "일반시험법",
                "일반정보",
            )
        ):
            return "생물학적제제 기준 및 시험방법"

        return "참고자료"

    def _split_rag_pdf_page(self, text: str, max_chars: int = 2600) -> list[str]:
        text = clean_text(text)
        if not text:
            return []

        paragraphs = [clean_text(p) for p in re.split(r"\n\s*\n+", text) if clean_text(p)]
        if not paragraphs:
            paragraphs = [text]

        chunks: list[str] = []
        current = ""
        for paragraph in paragraphs:
            if len(paragraph) > max_chars:
                if current:
                    chunks.append(current)
                    current = ""
                for start in range(0, len(paragraph), max_chars):
                    chunk = clean_text(paragraph[start : start + max_chars])
                    if chunk:
                        chunks.append(chunk)
                continue

            candidate = f"{current}\n{paragraph}".strip() if current else paragraph
            if len(candidate) > max_chars and current:
                chunks.append(current)
                current = paragraph
            else:
                current = candidate

        if current:
            chunks.append(current)

        return chunks

    def _build_json_doc_text(self, item: dict) -> str:
        pieces = [
            item.get("kind"),
            item.get("code"),
            item.get("name"),
            item.get("symbol"),
            item.get("property"),
            item.get("ref_unit"),
            item.get("ref_value"),
            item.get("text"),
        ]
        return self._normalize_for_search(" ".join(str(x) for x in pieces if x not in (None, "")))

    def _build_jsonl_doc_text(self, item: dict) -> str:
        fields = [
            item.get("major_category_ko"),
            item.get("property"),
            item.get("display_name_ko"),
            item.get("display_name_en"),
            item.get("canonical_ucum"),
            item.get("canonical_ascii"),
            item.get("canonical_symbol"),
            item.get("display_symbol_ui"),
            " ".join(item.get("display_symbol_variants", []) or []),
            " ".join(item.get("aliases", []) or []),
            item.get("text"),
        ]
        return self._normalize_for_search(" ".join(str(x) for x in fields if x not in (None, "")))

    def _load_json_file(self, path: Path, docs: list[RagDoc], seen: set[tuple[str, str]]) -> None:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)

        for item in data:
            text = self._build_json_doc_text(item)
            if not text:
                continue
            key = (str(path.resolve()), item.get("code") or item.get("text") or text)
            if key in seen:
                continue
            seen.add(key)
            meta = dict(item)
            meta["rag_source_file"] = path.name
            meta["rag_source_type"] = "json"
            meta["rag_domain"] = "기호 사전"
            docs.append(RagDoc(idx=len(docs), text=text, meta=meta))

        self._append_loaded_source(path.name)

    def _load_jsonl_file(self, path: Path, docs: list[RagDoc], seen: set[tuple[str, str]]) -> None:
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                item = json.loads(line)
                text = self._build_jsonl_doc_text(item)
                if not text:
                    continue
                key = ("jsonl", item.get("canonical_ucum") or item.get("id") or text)
                if key in seen:
                    continue
                seen.add(key)
                meta = dict(item)
                meta["rag_source_file"] = path.name
                meta["rag_source_type"] = "jsonl"
                meta["rag_domain"] = "기호 사전"
                docs.append(RagDoc(idx=len(docs), text=text, meta=meta))

        self._append_loaded_source(path.name)

    def _load_xlsx_file(self, path: Path, docs: list[RagDoc], seen: set[tuple[str, str]]) -> None:
        wb = load_workbook(path, data_only=True)
        if not wb.sheetnames:
            return
        ws = wb[wb.sheetnames[0]]

        for row in ws.iter_rows(min_row=3, values_only=True):
            (
                row_no,
                ucum_code,
                description,
                comment,
                last_updated,
                version_correction,
                corrected_by,
                previous_row_no,
                previous_ucum_version,
                change_description,
                *_,
            ) = row + (None,) * max(0, 10 - len(row))

            if not ucum_code:
                continue

            text = self._normalize_for_search(
                " ".join(
                    str(x)
                    for x in [
                        ucum_code,
                        description,
                        comment,
                        previous_ucum_version,
                        change_description,
                        version_correction,
                    ]
                    if x not in (None, "")
                )
            )
            if not text:
                continue

            key = ("xlsx", str(ucum_code))
            if key in seen:
                continue
            seen.add(key)
            meta = {
                "row_no": row_no,
                "ucum_code": ucum_code,
                "description": description,
                "comment": comment,
                "last_updated": str(last_updated) if last_updated else None,
                "version_correction": version_correction,
                "corrected_by": corrected_by,
                "previous_row_no": previous_row_no,
                "previous_ucum_version": previous_ucum_version,
                "change_description": change_description,
                "rag_source_file": path.name,
                "rag_source_type": "xlsx",
                "rag_domain": "기호 사전",
            }
            docs.append(RagDoc(idx=len(docs), text=text, meta=meta))

        self._append_loaded_source(path.name)

    def _load_rag_data_pdfs(self, docs: list[RagDoc], seen: set[tuple[str, str]]) -> None:
        if fitz is None or not RAG_DATA_DIR.exists():
            return

        for path in sorted(RAG_DATA_DIR.glob("*.pdf")):
            if not path.exists():
                continue

            domain = self._rag_domain_for_path(path)

            try:
                pdf = fitz.open(path)
            except Exception:
                continue

            try:
                for page_index in range(pdf.page_count):
                    try:
                        raw_page_text = pdf.load_page(page_index).get_text("text")
                    except Exception:
                        continue

                    for chunk_index, chunk in enumerate(self._split_rag_pdf_page(raw_page_text)):
                        text = self._normalize_for_search(chunk)
                        if len(text) < 20:
                            continue

                        key = (str(path.resolve()), f"{page_index + 1}:{chunk_index}:{text[:160]}")
                        if key in seen:
                            continue
                        seen.add(key)

                        meta = {
                            "rag_source_file": path.name,
                            "rag_source_type": "pdf",
                            "rag_domain": domain,
                            "page": page_index + 1,
                            "chunk": chunk_index + 1,
                        }
                        docs.append(RagDoc(idx=len(docs), text=text, meta=meta))
            finally:
                pdf.close()

            self._append_loaded_source(path.name)

    def _load(self) -> None:
        docs: list[RagDoc] = []
        seen: set[tuple[str, str]] = set()

        for path in self._existing_paths(UCUM_JSON_CANDIDATES):
            self._load_json_file(path, docs, seen)

        for path in self._existing_paths(UCUM_JSONL_CANDIDATES):
            self._load_jsonl_file(path, docs, seen)

        for path in self._existing_paths(UCUM_XLSX_CANDIDATES):
            self._load_xlsx_file(path, docs, seen)

        self._load_rag_data_pdfs(docs, seen)

        if not docs:
            raise FileNotFoundError(
                "UCUM RAG 소스를 찾을 수 없습니다. ucum_rag_docs.json / ucum_rag_image_units_ui_micro.jsonl / TableOfExampleUcumCodesForElectronicMessaging.xlsx 중 하나 이상이 필요합니다."
            )

        self.docs = docs
        corpus = [d.text for d in docs]

        self.word_vectorizer = TfidfVectorizer(
            ngram_range=(1, 2),
            lowercase=False,
            token_pattern=r"(?u)[\w\[\]/.%'*-]+",
        )
        self.char_vectorizer = TfidfVectorizer(
            analyzer="char_wb",
            ngram_range=(2, 5),
            lowercase=False,
        )
        self.word_matrix = self.word_vectorizer.fit_transform(corpus)
        self.char_matrix = self.char_vectorizer.fit_transform(corpus)

    def search(self, query: str, top_k: int = 5) -> list[RagDoc]:
        query = self._normalize_for_search(query)
        if (
            not query
            or not self.docs
            or self.word_vectorizer is None
            or self.char_vectorizer is None
            or self.word_matrix is None
            or self.char_matrix is None
        ):
            return []

        word_query = self.word_vectorizer.transform([query])
        char_query = self.char_vectorizer.transform([query])
        word_scores = linear_kernel(word_query, self.word_matrix).flatten()
        char_scores = linear_kernel(char_query, self.char_matrix).flatten()
        scores = (0.65 * word_scores) + (0.35 * char_scores)

        ranked = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)
        out: list[RagDoc] = []
        for i in ranked[:top_k]:
            if scores[i] <= 0:
                continue
            out.append(self.docs[i])
        return out
