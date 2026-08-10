from __future__ import annotations

import hashlib
import csv
import json
import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import FAIL_LABEL, HOLD_LABEL, PASS_LABEL
from .utils import clean_text


SUMMARY_SHEET = "회귀요약"
DETAIL_SHEET = "판정상세"
CATALOG_SHEET = "규칙카탈로그"
CONTEXT_SHEET = "RAG_MD사용내역"
HEADER_ROW = 4
GOLD_ROW = 5
FIRST_RUN_ROW = 6


@dataclass
class StatusCounts:
    passed: int = 0
    failed: int = 0
    held: int = 0

    @property
    def total(self) -> int:
        return self.passed + self.failed + self.held

    def add(self, status: str) -> None:
        normalized = _normalize_status(status)
        if normalized == "pass":
            self.passed += 1
        elif normalized == "fail":
            self.failed += 1
        elif normalized == "hold":
            self.held += 1

    def compact(self) -> str:
        return f"{self.passed}/{self.failed}/{self.held}"


@dataclass
class RuleObservation:
    rule_id: str
    display_name: str
    executor: str
    counts: StatusCounts = field(default_factory=StatusCounts)
    implementation_status: str = "구현"
    note: str = ""

    def compact(self) -> str:
        if self.implementation_status != "구현":
            return self.implementation_status
        return self.counts.compact()


@dataclass
class RuleCatalogEntry:
    rule_id: str
    section: str
    text: str
    line_number: int
    executor: str
    implementation_status: str


@dataclass
class RegressionLogResult:
    workbook_path: Path
    run_id: str
    appended: bool
    gold_status: str


RULE_COLUMNS = [
    ("test_result", "시험기준/결과 비교", "evaluation_status"),
    ("page_number", "페이지 번호 검증", "page_number_validator"),
    ("manufacturing_no", "제조번호 비교", "manufacturing_info_validator"),
    ("manufacturing_date", "제조일자 비교", "manufacturing_info_validator"),
    ("manufacturing_quantity", "제조량 비교", "manufacturing_info_validator"),
    ("test_date_window", "시험일 허용구간", "manufacturing_info_validator"),
    ("process_order", "날짜 선행관계 비교", "manufacturing_summary_validator"),
]


FIXED_HEADERS = [
    "기록구분",
    "정답지상태",
    "실행시각",
    "PDF파일",
    "회사",
    "제품",
    "LLM모델",
    "전체판정(충족/불충족/보류)",
    "정답지대비",
]


def _normalize_status(status: str) -> str:
    value = clean_text(status).lower().replace(" ", "")
    if value in {
        PASS_LABEL.lower(),
        "합격",
        "충족",
        "적합",
        "pass",
        "passed",
        "true",
    }:
        return "pass"
    if value in {
        FAIL_LABEL.lower(),
        "불합격",
        "불충족",
        "부적합",
        "fail",
        "failed",
        "false",
    }:
        return "fail"
    if value in {
        HOLD_LABEL.lower(),
        "보류",
        "확인필요",
        "확인필요함",
        "판단불가",
        "hold",
        "pending",
        "unknown",
    }:
        return "hold"
    return ""


def _sanitize_filename(value: str, fallback: str = "document") -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", clean_text(value))
    value = re.sub(r"\s+", " ", value).strip(" .")
    return (value or fallback)[:120]


@lru_cache(maxsize=128)
def _file_sha256_cached(path_text: str, size: int, mtime_ns: int) -> str:
    del size, mtime_ns
    digest = hashlib.sha256()
    with Path(path_text).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_sha256(path: Path) -> str:
    path = Path(path).resolve()
    stat = path.stat()
    return _file_sha256_cached(str(path), stat.st_size, stat.st_mtime_ns)


def _hash_named_parts(parts: Iterable[tuple[str, str]]) -> str:
    digest = hashlib.sha256()
    for name, value in sorted(parts):
        digest.update(name.encode("utf-8", errors="replace"))
        digest.update(b"\0")
        digest.update(value.encode("utf-8", errors="replace"))
        digest.update(b"\0")
    return digest.hexdigest()


def rule_code_fingerprint(project_root: Path) -> str:
    project_root = Path(project_root)
    candidates = [
        project_root / "sp_judgement_bridge.py",
        project_root / "sp_pdf_judger" / "judgement.py",
        project_root / "sp_pdf_judger" / "llm.py",
        project_root / "sp_pdf_judger" / "pipeline.py",
        project_root / "sp_pdf_judger" / "rag.py",
        project_root / "sp_pdf_judger" / "domain_details.py",
        project_root / "sp_pdf_judger" / "manufacturing_info_validator.py",
        project_root / "sp_pdf_judger" / "page_number_validator.py",
        project_root / "sp_pdf_judger" / "rule_regression.py",
    ]
    validator_dir = project_root / "sp_pdf_judger"
    if validator_dir.exists():
        candidates.extend(sorted(validator_dir.glob("*_validator.py")))

    parts: list[tuple[str, str]] = []
    for path in dict.fromkeys(candidates):
        if path.exists() and path.is_file():
            parts.append((path.relative_to(project_root).as_posix(), file_sha256(path)))
    return _hash_named_parts(parts)


def rag_fingerprint(project_root: Path, loaded_sources: Iterable[str]) -> str:
    project_root = Path(project_root)
    parts: list[tuple[str, str]] = []
    for source in sorted({clean_text(item) for item in loaded_sources if clean_text(item)}):
        path = Path(source)
        if not path.is_absolute():
            path = project_root / path
        if path.exists() and path.is_file():
            stat = path.stat()
            signature = f"{stat.st_size}:{stat.st_mtime_ns}"
        else:
            signature = "reported-source"
        parts.append((source, signature))
    return _hash_named_parts(parts) if parts else "none"


def _executor_for_rule(text: str) -> tuple[str, str]:
    normalized = re.sub(r"\s+", "", text).lower()
    if "페이지번호" in normalized:
        return "page_number_validator", "미구현"
    if "제조번호" in normalized:
        return "manufacturing_info_validator", "구현"
    if any(token in normalized for token in ("제조년월일", "제조일", "분병일자", "포장일자")):
        return "manufacturing_info_validator", "구현"
    if any(token in normalized for token in ("제조량", "사용량", "분량", "투입량")):
        return "manufacturing_info_validator", "구현"
    if any(token in normalized for token in ("선후관계", "병렬공정", "합류")):
        return "manufacturing_summary_validator", "구현"
    if any(token in normalized for token in ("시험기준", "시험결과", "수치기준", "범위기준", "정성기준", "복수조건", "용기규격", "농도")):
        return "evaluation_status", "구현"
    if any(token in normalized for token in ("약전", "허가서", "참고문서", "검색근거", "근거")):
        return "rag_permit_trace", "관측"
    if "보류" in normalized:
        return "evaluation_status", "구현"
    if any(token in normalized for token in ("ocr", "행과열", "중복집계", "추적가능", "목차번호")):
        return "structure_validator", "미구현"
    return "unmapped", "미구현"


def parse_overall_rule_catalog(overall_md_path: Path) -> list[RuleCatalogEntry]:
    overall_md_path = Path(overall_md_path)
    if not overall_md_path.exists():
        return []

    section = ""
    entries: list[RuleCatalogEntry] = []
    in_frontmatter = False
    for line_number, raw_line in enumerate(
        overall_md_path.read_text(encoding="utf-8").splitlines(),
        start=1,
    ):
        stripped = raw_line.strip()
        if line_number == 1 and stripped == "---":
            in_frontmatter = True
            continue
        if in_frontmatter:
            if stripped == "---":
                in_frontmatter = False
            continue
        if stripped.startswith("#"):
            section = stripped.lstrip("#").strip()
            continue
        if not stripped.startswith("- "):
            continue

        text = stripped[2:].strip()
        executor, implementation_status = _executor_for_rule(text)
        digest = hashlib.sha1(f"{section}\n{text}".encode("utf-8")).hexdigest()[:12]
        entries.append(
            RuleCatalogEntry(
                rule_id=f"overall-{digest}",
                section=section,
                text=text,
                line_number=line_number,
                executor=executor,
                implementation_status=implementation_status,
            )
        )
    return entries


def _iter_evaluation_statuses(result: Any) -> Iterable[tuple[str, str]]:
    for evaluation in list(getattr(result, "evaluations", []) or []):
        source = clean_text(getattr(evaluation, "source", ""))
        lot_judgements = list(getattr(evaluation, "lot_judgements", []) or [])
        if lot_judgements:
            for item in lot_judgements:
                if isinstance(item, dict):
                    yield clean_text(item.get("status", "")), clean_text(item.get("source", source))
            continue
        yield clean_text(getattr(evaluation, "final_status", "")), source


def _manufacturing_observations(result: Any) -> dict[str, RuleObservation]:
    observations = {
        "manufacturing_no": RuleObservation(
            "manufacturing_no",
            "제조번호 비교",
            "manufacturing_info_validator",
        ),
        "manufacturing_date": RuleObservation(
            "manufacturing_date",
            "제조일자 비교",
            "manufacturing_info_validator",
        ),
        "manufacturing_quantity": RuleObservation(
            "manufacturing_quantity",
            "제조량 비교",
            "manufacturing_info_validator",
        ),
        "test_date_window": RuleObservation(
            "test_date_window",
            "시험일 허용구간",
            "manufacturing_info_validator",
        ),
    }
    try:
        from .manufacturing_info_validator import validate_manufacturing_info_consistency

        validation_results = validate_manufacturing_info_consistency(result)
    except Exception as exc:
        for observation in observations.values():
            observation.implementation_status = "실행오류"
            observation.note = f"{type(exc).__name__}: {exc}"
        return observations

    for validation in validation_results:
        for field_result in list(getattr(validation, "fields", []) or []):
            field_name = clean_text(getattr(field_result, "field_name", ""))
            status = clean_text(getattr(field_result, "status", ""))
            if "제조번호" in field_name:
                observations["manufacturing_no"].counts.add(status)
            elif any(token in field_name for token in ("제조년월일", "제조일자", "제조일")):
                observations["manufacturing_date"].counts.add(status)
            elif any(token in field_name for token in ("제조량", "제조수량")):
                observations["manufacturing_quantity"].counts.add(status)
        for test_date in list(getattr(validation, "test_dates", []) or []):
            observations["test_date_window"].counts.add(
                clean_text(getattr(test_date, "status", ""))
            )

    for observation in observations.values():
        if observation.counts.total == 0:
            observation.note = "현재 문서에서 해당 비교 결과가 생성되지 않았습니다."
    return observations


def collect_rule_observations(result: Any) -> list[RuleObservation]:
    evaluation_counts = StatusCounts()
    reference_counts = StatusCounts()
    hold_counts = StatusCounts()
    sources: list[str] = []

    for status, source in _iter_evaluation_statuses(result):
        evaluation_counts.add(status)
        normalized = _normalize_status(status)
        if normalized == "hold":
            hold_counts.add(status)
        source_lower = source.lower()
        if any(token in source_lower for token in ("rag", "permit", "허가", "gemini", "llm")):
            reference_counts.add(status)
            if source:
                sources.append(source)

    records = list(getattr(result, "extracted_records", []) or [])
    ocr_suspect_count = sum(bool(getattr(record, "ocr_suspect", False)) for record in records)

    observations = {
        "test_result": RuleObservation(
            "test_result",
            "시험기준/결과 비교",
            "evaluation_status",
            counts=evaluation_counts,
            note="표 시험은 lot_judgements의 각 행을 1건으로 집계합니다.",
        ),
        "document_structure": RuleObservation(
            "document_structure",
            "문서 구조/추적성",
            "structure_validator",
            implementation_status="미구현",
            note=f"OCR 의심 레코드 관측 {ocr_suspect_count}건. 구조 전체를 합격/불합격으로 판정하는 전용 검증기는 아직 없습니다.",
        ),
        "page_number": RuleObservation(
            "page_number",
            "페이지 번호 검증",
            "page_number_validator",
            implementation_status="미구현",
            note="overall.md에는 규칙이 있으나 원본 코드에 page_number_validator.py가 없습니다.",
        ),
        "process_order": RuleObservation(
            "process_order",
            "날짜 선행관계 비교",
            "manufacturing_summary_validator",
            note=clean_text(getattr(result, "manufacturing_summary_reason", "")),
        ),
        "reference_use": RuleObservation(
            "reference_use",
            "참고 문서 적용",
            "rag_permit_trace",
            counts=reference_counts,
            implementation_status="구현",
            note="판정 source에 RAG, permit, 허가서, Gemini 또는 LLM이 명시된 결과만 집계합니다."
            + (f" 사용 source: {', '.join(sorted(set(sources)))}" if sources else ""),
        ),
        "hold_policy": RuleObservation(
            "hold_policy",
            "보류 원칙",
            "evaluation_status",
            counts=hold_counts,
            note="전체 시험 판정 중 최종 보류만 집계하며 충족/불충족 값은 0으로 둡니다.",
        ),
    }

    process_status = clean_text(getattr(result, "manufacturing_summary_status", ""))
    if process_status:
        observations["process_order"].counts.add(process_status)
    else:
        observations["process_order"].implementation_status = "미실행"

    observations.update(_manufacturing_observations(result))
    return [
        observations[rule_id]
        for rule_id, _, _ in RULE_COLUMNS
    ]


def _summary_counts(result: Any) -> StatusCounts:
    summary = getattr(result, "summary", None)
    if summary is not None:
        return StatusCounts(
            passed=int(getattr(summary, "passed", 0) or 0),
            failed=int(getattr(summary, "failed", 0) or 0),
            held=int(getattr(summary, "held", 0) or 0),
        )
    counts = StatusCounts()
    for status, _ in _iter_evaluation_statuses(result):
        counts.add(status)
    return counts


def _summary_counts_from_csv(path: Path | str | None) -> StatusCounts | None:
    if not path:
        return None
    summary_path = Path(path)
    if not summary_path.exists():
        return None
    try:
        with summary_path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    except (OSError, UnicodeError, csv.Error):
        return None
    if not rows:
        return None

    total_row = next(
        (
            row
            for row in rows
            if clean_text(row.get("제조명", "")) == "전체"
        ),
        None,
    )
    if total_row is None:
        return None

    def _integer(column: str) -> int:
        value = clean_text(total_row.get(column, "")).replace(",", "")
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return 0

    return StatusCounts(
        passed=_integer("검수합격"),
        failed=_integer("검수불합격"),
        held=_integer("검수보류"),
    )


def _llm_call_explanation(metadata: dict[str, Any], raw_counts: StatusCounts) -> str:
    call_count = int(metadata.get("llm_call_count", 0) or 0)
    success_count = int(metadata.get("llm_success_count", 0) or 0)
    last_error = clean_text(metadata.get("llm_last_error", ""))
    if call_count:
        explanation = f"보조 LLM {call_count}건 호출, {success_count}건 성공"
        if last_error:
            explanation += f"; 마지막 오류: {last_error}"
        return explanation
    if last_error:
        return f"LLM 호출 0건; 초기화 또는 호출 오류: {last_error}"
    if raw_counts.total and raw_counts.held == 0:
        return (
            f"직접 판정 규칙이 전체 시험행 {raw_counts.total}건을 모두 확정해 "
            "보조 LLM 호출 대상이 없었습니다."
        )
    if not bool(metadata.get("llm_enabled", False)):
        return "LLM이 비활성화되어 호출하지 않았습니다."
    return "LLM 보조판정 호출 조건에 해당한 미확정 항목이 없었습니다."


def _run_id(
    *,
    pdf_sha256: str,
    permit_hashes: list[str],
    rule_fingerprint: str,
    rag_fp: str,
    detail_fingerprint: str,
    llm_model: str,
    llm_execution_signature: str,
) -> str:
    return _hash_named_parts(
        [
            ("pdf", pdf_sha256),
            ("permits", "|".join(sorted(permit_hashes))),
            ("rules", rule_fingerprint),
            ("rag", rag_fp),
            ("details", detail_fingerprint),
            ("llm", llm_model),
            ("llm_execution", llm_execution_signature),
        ]
    )[:24]


def _header_map(sheet) -> dict[str, int]:
    return {
        clean_text(sheet.cell(HEADER_ROW, column).value): column
        for column in range(1, sheet.max_column + 1)
        if clean_text(sheet.cell(HEADER_ROW, column).value)
    }


def _gold_comparison_formula(sheet, row: int) -> str:
    headers = _header_map(sheet)
    comparison_headers = [
        FIXED_HEADERS[7],
        *[display_name for _, display_name, _ in RULE_COLUMNS],
    ]
    labels = [
        "\uc804\uccb4\ud310\uc815",
        *[display_name for _, display_name, _ in RULE_COLUMNS],
    ]
    configured_terms: list[str] = []
    mismatch_count_terms: list[str] = []
    mismatch_terms: list[str] = []
    unset = "\ubbf8\uc124\uc815"

    for header, label in zip(comparison_headers, labels):
        column = headers.get(header)
        if not column:
            continue
        letter = get_column_letter(column)
        gold_ref = f"${letter}${GOLD_ROW}"
        run_ref = f"{letter}{row}"
        safe_label = label.replace('"', '""')
        configured_terms.append(
            f'IF(OR({gold_ref}="",{gold_ref}="{unset}"),0,1)'
        )
        mismatch_count_terms.append(
            f'IF(OR({gold_ref}="",{gold_ref}="{unset}",'
            f'{gold_ref}={run_ref}),0,1)'
        )
        mismatch_terms.append(
            f'IF(OR({gold_ref}="",{gold_ref}="{unset}",'
            f'{gold_ref}={run_ref}),"","[{safe_label}]")'
        )

    configured = "+".join(configured_terms) or "0"
    mismatch_count = "+".join(mismatch_count_terms) or "0"
    mismatches = "&".join(mismatch_terms) or '""'
    return (
        f'=IF(({configured})=0,'
        '"\ubbf8\uc124\uc815: \uc815\ub2f5\uc9c0\uac00 \uc124\uc815\ub418\uc9c0 \uc54a\uc558\uc2b5\ub2c8\ub2e4.",'
        f'IF(({mismatch_count})=0,'
        '"\uc77c\uce58: \uc815\ub2f5\uc9c0\uc640 \ubaa8\ub4e0 \uc124\uc815 \ud56d\ubaa9\uc774 \uc77c\uce58\ud569\ub2c8\ub2e4.",'
        f'"\ubd88\uc77c\uce58: "&{mismatches}))'
    )


def _refresh_gold_comparison_formulas(sheet) -> bool:
    headers = _header_map(sheet)
    comparison_column = headers.get(FIXED_HEADERS[8])
    if not comparison_column:
        return False

    changed = False
    for row in range(FIRST_RUN_ROW, sheet.max_row + 1):
        if not any(
            clean_text(sheet.cell(row, column).value)
            for column in range(1, sheet.max_column + 1)
        ):
            continue
        formula = _gold_comparison_formula(sheet, row)
        cell = sheet.cell(row, comparison_column)
        if cell.value != formula:
            cell.value = formula
            changed = True
    return changed


def _ensure_summary_sheet(workbook, observation_names: list[str]) -> tuple[Any, bool]:
    if SUMMARY_SHEET in workbook.sheetnames:
        sheet = workbook[SUMMARY_SHEET]
    else:
        sheet = workbook.active
        sheet.title = SUMMARY_SHEET
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        sheet.cell(1, 1, "PDF별 전체 규칙 회귀 로그")
        sheet.cell(2, 1, "규칙 결과 셀 형식: 충족/불충족/보류. 정답지 행은 자동으로 덮어쓰지 않습니다.")

    headers = FIXED_HEADERS + observation_names
    existing = _header_map(sheet)
    old_headers = [
        clean_text(sheet.cell(HEADER_ROW, column).value)
        for column in range(1, sheet.max_column + 1)
        if clean_text(sheet.cell(HEADER_ROW, column).value)
    ]
    old_rows: list[dict[str, Any]] = []
    if existing:
        for row in range(GOLD_ROW, sheet.max_row + 1):
            values = {
                header: sheet.cell(row, column).value
                for header, column in existing.items()
            }
            if row == GOLD_ROW or any(clean_text(value) for value in values.values()):
                old_rows.append(values)

    old_gold = old_rows[0] if old_rows else {}
    old_gold_status = clean_text(old_gold.get("정답지상태", ""))
    legacy_auto_draft = old_gold_status in {"초안", "초안(첫 실행값)"}
    result_headers = {"전체판정(충족/불충족/보류)", *observation_names}
    layout_changed = old_headers != headers or legacy_auto_draft

    if sheet.max_row >= HEADER_ROW:
        sheet.delete_rows(HEADER_ROW, sheet.max_row - HEADER_ROW + 1)
    if sheet.max_column > len(headers):
        sheet.delete_cols(len(headers) + 1, sheet.max_column - len(headers))

    for column, header in enumerate(headers, start=1):
        sheet.cell(HEADER_ROW, column, header)

    gold_values = {
        header: old_gold.get(header)
        for header in headers
    }
    gold_values["기록구분"] = "정답지"
    if legacy_auto_draft or not old_gold:
        gold_values["정답지상태"] = "입력 대기"
        for header in result_headers:
            gold_values[header] = None
    elif not clean_text(gold_values.get("정답지상태")):
        gold_values["정답지상태"] = "입력 대기"

    for column, header in enumerate(headers, start=1):
        sheet.cell(GOLD_ROW, column, gold_values.get(header))

    for old_row in old_rows[1:]:
        row = sheet.max_row + 1
        for column, header in enumerate(headers, start=1):
            sheet.cell(row, column, old_row.get(header))

    formula_changed = _refresh_gold_comparison_formulas(sheet)
    return sheet, layout_changed or formula_changed


def _style_workbook(workbook) -> None:
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    gold_fill = PatternFill("solid", fgColor="FFF2CC")
    title_font = Font(size=16, bold=True, color="17365D")
    calculation = getattr(workbook, "calculation", None)
    if calculation is not None:
        calculation.calcMode = "auto"
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True

    summary = workbook[SUMMARY_SHEET]
    summary.cell(1, 1).font = title_font
    summary.freeze_panes = f"A{FIRST_RUN_ROW}"
    summary.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(summary.max_column)}{summary.max_row}"
    for cell in summary[HEADER_ROW]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in summary[GOLD_ROW]:
        cell.fill = gold_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in summary.iter_rows(min_row=FIRST_RUN_ROW):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 12,
        "B": 22,
        "C": 20,
        "D": 42,
        "E": 20,
        "F": 24,
        "G": 20,
        "H": 20,
        "I": 20,
        "J": 18,
        "K": 24,
        "L": 24,
        "M": 12,
        "N": 12,
        "O": 24,
        "P": 32,
    }
    for column, width in widths.items():
        summary.column_dimensions[column].width = width
    for column in range(len(FIXED_HEADERS) + 1, summary.max_column + 1):
        summary.column_dimensions[get_column_letter(column)].width = 24
        summary.cell(HEADER_ROW, column).comment = Comment(
            "값 순서: 충족/불충족/보류",
            "SP AI",
        )

    for sheet_name in (DETAIL_SHEET, CATALOG_SHEET, CONTEXT_SHEET):
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [
                len(clean_text(sheet.cell(row, column).value))
                for row in range(1, min(sheet.max_row, 40) + 1)
            ]
            sheet.column_dimensions[get_column_letter(column)].width = min(
                max(max(values, default=10) + 2, 12),
                70,
            )


def _ensure_tabular_sheet(workbook, name: str, headers: list[str]) -> Any:
    if name in workbook.sheetnames:
        sheet = workbook[name]
    else:
        sheet = workbook.create_sheet(name)
    if sheet.max_row == 1 and not clean_text(sheet.cell(1, 1).value):
        for column, header in enumerate(headers, start=1):
            sheet.cell(1, column, header)
    return sheet


def _has_run(workbook, run_id: str) -> bool:
    for sheet_name in (CONTEXT_SHEET, DETAIL_SHEET, CATALOG_SHEET):
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        run_column = next(
            (
                column
                for column in range(1, sheet.max_column + 1)
                if clean_text(sheet.cell(1, column).value) == "실행_ID"
            ),
            None,
        )
        if run_column and any(
            clean_text(sheet.cell(row, run_column).value) == run_id
            for row in range(2, sheet.max_row + 1)
        ):
            return True
    return False


def _gold_comparison(
    sheet,
    observations: list[RuleObservation],
    overall: str,
) -> tuple[str, str]:
    headers = _header_map(sheet)
    mismatches: list[str] = []
    compared = 0

    expected_overall = clean_text(
        sheet.cell(GOLD_ROW, headers["전체판정(충족/불충족/보류)"]).value
    )
    if expected_overall and expected_overall != "미설정":
        compared += 1
        if expected_overall != overall:
            mismatches.append("전체판정")

    for observation in observations:
        column = headers[observation.display_name]
        expected = clean_text(sheet.cell(GOLD_ROW, column).value)
        if not expected or expected == "미설정":
            continue
        compared += 1
        if expected != observation.compact():
            mismatches.append(observation.display_name)

    if not compared:
        return "미설정", "정답지가 설정되지 않았습니다."
    if mismatches:
        return "불일치", ", ".join(mismatches)
    return "일치", "정답지와 모든 설정 항목이 일치합니다."


def _append_catalog_rows(
    workbook,
    *,
    run_id: str,
    detail_fingerprint: str,
    entries: list[RuleCatalogEntry],
) -> None:
    sheet = _ensure_tabular_sheet(
        workbook,
        CATALOG_SHEET,
        [
            "실행_ID",
            "MD지문",
            "규칙_ID",
            "섹션",
            "MD규칙문장",
            "원문줄",
            "실행기",
            "구현상태",
        ],
    )
    existing = {
        (clean_text(sheet.cell(row, 1).value), clean_text(sheet.cell(row, 3).value))
        for row in range(2, sheet.max_row + 1)
    }
    for entry in entries:
        if (run_id, entry.rule_id) in existing:
            continue
        sheet.append(
            [
                run_id,
                detail_fingerprint,
                entry.rule_id,
                entry.section,
                entry.text,
                entry.line_number,
                entry.executor,
                entry.implementation_status,
            ]
        )


def _append_detail_rows(
    workbook,
    *,
    run_id: str,
    timestamp: str,
    observations: list[RuleObservation],
) -> None:
    sheet = _ensure_tabular_sheet(
        workbook,
        DETAIL_SHEET,
        [
            "실행_ID",
            "실행시각",
            "규칙_ID",
            "규칙명",
            "충족",
            "불충족",
            "보류",
            "합계",
            "구현상태",
            "실행기",
            "비고",
        ],
    )
    existing = {
        (clean_text(sheet.cell(row, 1).value), clean_text(sheet.cell(row, 3).value))
        for row in range(2, sheet.max_row + 1)
    }
    for observation in observations:
        if (run_id, observation.rule_id) in existing:
            continue
        sheet.append(
            [
                run_id,
                timestamp,
                observation.rule_id,
                observation.display_name,
                observation.counts.passed,
                observation.counts.failed,
                observation.counts.held,
                observation.counts.total,
                observation.implementation_status,
                observation.executor,
                observation.note,
            ]
        )


def _append_context_row(
    workbook,
    *,
    run_id: str,
    timestamp: str,
    metadata: dict[str, Any],
    permit_paths: list[Path],
) -> None:
    sheet = _ensure_tabular_sheet(
        workbook,
        CONTEXT_SHEET,
        [
            "실행_ID",
            "실행시각",
            "RAG문서수",
            "RAG소스",
            "연결허가서",
            "MD소스",
            "매칭회사",
            "매칭제품",
            "LLM전송_MD컨텍스트",
            "LLM마지막오류",
        ],
    )
    if any(clean_text(sheet.cell(row, 1).value) == run_id for row in range(2, sheet.max_row + 1)):
        return
    context = clean_text(metadata.get("domain_detail_context", ""))
    if len(context) > 32000:
        context = context[:31970] + "\n...[Excel 셀 길이 제한으로 생략]"
    sheet.append(
        [
            run_id,
            timestamp,
            int(metadata.get("rag_doc_count", 0) or 0),
            "\n".join(str(item) for item in metadata.get("rag_sources", []) or []),
            "\n".join(str(path) for path in permit_paths),
            "\n".join(str(item) for item in metadata.get("domain_detail_sources", []) or []),
            clean_text(metadata.get("domain_detail_matched_company", "")),
            clean_text(metadata.get("domain_detail_matched_product", "")),
            context,
            clean_text(metadata.get("llm_last_error", "")),
        ]
    )


def _atomic_save(workbook, workbook_path: Path) -> None:
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = workbook_path.with_name(f".{workbook_path.stem}.tmp.xlsx")
    try:
        workbook.save(temp_path)
        temp_path.replace(workbook_path)
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def append_rule_regression_log(
    result: Any,
    *,
    project_root: Path,
    permit_paths: Iterable[Path] = (),
    summary_counts_path: Path | str | None = None,
    output_dir: Path | None = None,
    rule_fingerprint_override: str = "",
) -> RegressionLogResult:
    if os.getenv("SP_RULE_REGRESSION_LOG", "1").strip().lower() in {"0", "false", "off", "no"}:
        pdf_path = Path(getattr(result, "pdf_path"))
        return RegressionLogResult(
            workbook_path=(output_dir or Path(project_root) / "OUTPUT") / f"{_sanitize_filename(pdf_path.stem)}.xlsx",
            run_id="disabled",
            appended=False,
            gold_status="비활성",
        )

    project_root = Path(project_root).resolve()
    pdf_path = Path(getattr(result, "pdf_path")).resolve()
    permit_paths = [Path(path).resolve() for path in permit_paths if Path(path).exists()]
    metadata = getattr(result, "metadata", {})
    metadata = metadata if isinstance(metadata, dict) else {}

    pdf_hash = file_sha256(pdf_path)
    permit_hashes = [file_sha256(path) for path in permit_paths]
    rule_fp = rule_fingerprint_override or rule_code_fingerprint(project_root)
    rag_fp = rag_fingerprint(project_root, metadata.get("rag_sources", []) or [])
    detail_fp = clean_text(metadata.get("domain_detail_fingerprint", "")) or "none"
    llm_model = clean_text(metadata.get("llm_model", "")) or "disabled"
    llm_execution_signature = ":".join(
        [
            str(int(metadata.get("llm_call_count", 0) or 0)),
            str(int(metadata.get("llm_success_count", 0) or 0)),
            hashlib.sha256(
                clean_text(metadata.get("llm_last_error", "")).encode(
                    "utf-8",
                    errors="replace",
                )
            ).hexdigest()[:12],
        ]
    )
    run_id = _run_id(
        pdf_sha256=pdf_hash,
        permit_hashes=permit_hashes,
        rule_fingerprint=rule_fp,
        rag_fp=rag_fp,
        detail_fingerprint=detail_fp,
        llm_model=llm_model,
        llm_execution_signature=llm_execution_signature,
    )

    output_dir = Path(output_dir or project_root / "OUTPUT")
    workbook_path = output_dir / (
        f"{_sanitize_filename(pdf_path.stem)}__{pdf_hash[:12]}__rule_regression.xlsx"
    )

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = Workbook()
    observation_names = [display_name for _, display_name, _ in RULE_COLUMNS]
    summary_sheet, layout_changed = _ensure_summary_sheet(workbook, observation_names)
    if _has_run(workbook, run_id):
        if layout_changed:
            _style_workbook(workbook)
            _atomic_save(workbook, workbook_path)
        gold_status = clean_text(
            summary_sheet.cell(GOLD_ROW, _header_map(summary_sheet)["정답지상태"]).value
        )
        return RegressionLogResult(
            workbook_path=workbook_path,
            run_id=run_id,
            appended=False,
            gold_status=gold_status,
        )

    observations = collect_rule_observations(result)
    raw_overall_counts = _summary_counts(result)
    display_overall_counts = (
        _summary_counts_from_csv(summary_counts_path)
        or raw_overall_counts
    )
    overall = display_overall_counts.compact()
    headers = _header_map(summary_sheet)

    comparison, comparison_note = _gold_comparison(summary_sheet, observations, overall)
    timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
    row = max(summary_sheet.max_row + 1, FIRST_RUN_ROW)
    values = {
        "기록구분": "실행결과",
        "정답지상태": "",
        "실행시각": timestamp,
        "PDF파일": pdf_path.name,
        "회사": clean_text(metadata.get("document_company", "")),
        "제품": clean_text(metadata.get("document_product", "")),
        "LLM모델": llm_model,
        "전체판정(충족/불충족/보류)": overall,
        "정답지대비": f"{comparison}: {comparison_note}",
    }
    values.update({item.display_name: item.compact() for item in observations})
    for header, value in values.items():
        summary_sheet.cell(row, headers[header], value)
    summary_sheet.cell(
        row,
        headers[FIXED_HEADERS[8]],
        _gold_comparison_formula(summary_sheet, row),
    )

    overall_md_path = project_root / "sp_pdf_judger" / "domain_details" / "overall.md"
    catalog = parse_overall_rule_catalog(overall_md_path)
    _append_detail_rows(
        workbook,
        run_id=run_id,
        timestamp=timestamp,
        observations=observations,
    )
    _append_catalog_rows(
        workbook,
        run_id=run_id,
        detail_fingerprint=detail_fp,
        entries=catalog,
    )
    _append_context_row(
        workbook,
        run_id=run_id,
        timestamp=timestamp,
        metadata=metadata,
        permit_paths=permit_paths,
    )
    _style_workbook(workbook)
    _atomic_save(workbook, workbook_path)

    gold_status = clean_text(
        summary_sheet.cell(GOLD_ROW, headers["정답지상태"]).value
    )
    return RegressionLogResult(
        workbook_path=workbook_path,
        run_id=run_id,
        appended=True,
        gold_status=gold_status,
    )


def set_gold_from_latest(workbook_path: Path, *, status: str = "확정") -> None:
    workbook_path = Path(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = _header_map(sheet)
    if sheet.max_row < FIRST_RUN_ROW:
        raise ValueError("정답지로 복사할 실행 결과가 없습니다.")
    latest_row = sheet.max_row
    for header in ["전체판정(충족/불충족/보류)"] + [
        display_name for _, display_name, _ in RULE_COLUMNS
    ]:
        column = headers.get(header)
        if column:
            sheet.cell(GOLD_ROW, column, sheet.cell(latest_row, column).value)
    sheet.cell(GOLD_ROW, headers["정답지상태"], status)
    _style_workbook(workbook)
    _atomic_save(workbook, workbook_path)


def workbook_summary(workbook_path: Path) -> dict[str, Any]:
    workbook_path = Path(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[SUMMARY_SHEET]
    headers = _header_map(sheet)
    latest_run_id = ""
    if CONTEXT_SHEET in workbook.sheetnames:
        context_sheet = workbook[CONTEXT_SHEET]
        context_headers = {
            clean_text(context_sheet.cell(1, column).value): column
            for column in range(1, context_sheet.max_column + 1)
            if clean_text(context_sheet.cell(1, column).value)
        }
        run_column = context_headers.get("실행_ID")
        if run_column and context_sheet.max_row >= 2:
            latest_run_id = clean_text(
                context_sheet.cell(context_sheet.max_row, run_column).value
            )
    return {
        "path": str(workbook_path),
        "gold_status": clean_text(
            sheet.cell(GOLD_ROW, headers["정답지상태"]).value
        ),
        "run_count": max(0, sheet.max_row - FIRST_RUN_ROW + 1),
        "latest_run_id": latest_run_id,
    }


def write_output_index(output_dir: Path) -> Path:
    output_dir = Path(output_dir)
    rows = [
        workbook_summary(path)
        for path in sorted(output_dir.glob("*__rule_regression.xlsx"))
    ]
    index_path = output_dir / "rule_regression_index.json"
    index_path.write_text(
        json.dumps({"workbooks": rows}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return index_path
