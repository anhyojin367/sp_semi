from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from sp_pdf_judger.config import FAIL_LABEL, HOLD_LABEL, PASS_LABEL
from sp_pdf_judger.rule_regression import (
    FIRST_RUN_ROW,
    GOLD_ROW,
    RULE_COLUMNS,
    SUMMARY_SHEET,
    RuleObservation,
    StatusCounts,
    append_rule_regression_log,
    parse_overall_rule_catalog,
    set_gold_from_latest,
    workbook_summary,
)
from sp_pdf_judger.schemas import Evaluation, ProcessingResult, Summary


def _make_result(pdf_path: Path) -> ProcessingResult:
    return ProcessingResult(
        pdf_path=pdf_path,
        preview_image_path=pdf_path.with_suffix(".png"),
        extracted_records=[],
        evaluations=[
            Evaluation(
                test_name="시험 A",
                final_status=PASS_LABEL,
                source="rule",
            ),
            Evaluation(
                test_name="시험 B",
                final_status=FAIL_LABEL,
                source="rag",
            ),
            Evaluation(
                test_name="시험 C",
                final_status=HOLD_LABEL,
                source="llm",
            ),
        ],
        tree=[],
        summary=Summary(passed=1, failed=1, held=1, total=3),
        metadata={
            "llm_model": "test-model",
            "llm_call_count": 2,
            "llm_success_count": 2,
            "rag_sources": ["rag_data/reference.pdf"],
            "domain_detail_sources": [
                "sp_pdf_judger/domain_details/overall.md",
                "sp_pdf_judger/domain_details/companies/example.md",
            ],
            "domain_detail_fingerprint": "detail-fingerprint",
            "domain_detail_context": "# 전체 공통\n시험기준과 시험결과를 비교한다.",
            "domain_detail_matched_company": "예시회사",
            "domain_detail_matched_product": "예시제품",
            "document_company": "예시회사",
            "document_product": "예시제품",
        },
        manufacturing_summary_status=PASS_LABEL,
        manufacturing_summary_reason="공정 순서가 정상입니다.",
    )


def _fake_manufacturing_observations(_result) -> dict[str, RuleObservation]:
    observations: dict[str, RuleObservation] = {}
    counts_by_rule = {
        "manufacturing_no": StatusCounts(passed=2, failed=0, held=0),
        "manufacturing_date": StatusCounts(passed=1, failed=1, held=0),
        "manufacturing_quantity": StatusCounts(passed=1, failed=0, held=1),
        "test_date_window": StatusCounts(passed=3, failed=0, held=0),
    }
    columns = {rule_id: (display_name, executor) for rule_id, display_name, executor in RULE_COLUMNS}
    for rule_id, counts in counts_by_rule.items():
        display_name, executor = columns[rule_id]
        observations[rule_id] = RuleObservation(
            rule_id=rule_id,
            display_name=display_name,
            executor=executor,
            counts=counts,
        )
    return observations


def test_append_creates_blank_gold_row_and_deduplicates_same_run(tmp_path, monkeypatch):
    project_root = tmp_path / "project"
    (project_root / "sp_pdf_judger" / "domain_details").mkdir(parents=True)
    (project_root / "sp_pdf_judger" / "domain_details" / "overall.md").write_text(
        "# 전체 공통\n- 시험기준과 시험결과를 비교한다.\n",
        encoding="utf-8",
    )
    pdf_path = project_root / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nregression-test")
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )

    first = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v1",
    )
    duplicate = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v1",
    )

    assert first.appended is True
    assert duplicate.appended is False
    assert first.run_id == duplicate.run_id

    workbook = load_workbook(first.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
    }
    assert sheet.cell(4, 1).value == "기록구분"
    assert sheet.max_row == FIRST_RUN_ROW
    assert sheet.cell(GOLD_ROW, headers["정답지상태"]).value == "입력 대기"
    assert sheet.cell(GOLD_ROW, headers["전체판정(충족/불충족/보류)"]).value is None
    assert sheet.cell(GOLD_ROW, headers["시험기준/결과 비교"]).value is None
    assert sheet.cell(FIRST_RUN_ROW, headers["시험기준/결과 비교"]).value == "1/1/1"
    assert sheet.cell(FIRST_RUN_ROW, headers["제조일자 비교"]).value == "1/1/0"
    removed_headers = {
        "PDF_SHA256",
        "실행_ID",
        "규칙코드지문",
        "RAG지문",
        "MD지문",
        "LLM호출",
        "LLM성공",
        "문서 구조/추적성",
        "참고 문서 적용",
        "보류 원칙",
        "LLM호출설명",
        "화면집계출처",
        "제조단계 미매핑 시험행",
    }
    assert removed_headers.isdisjoint(headers)


def test_summary_csv_is_logged_as_ui_total_without_hiding_raw_test_rows(
    tmp_path,
    monkeypatch,
):
    project_root = tmp_path / "project"
    (project_root / "sp_pdf_judger" / "domain_details").mkdir(parents=True)
    (project_root / "sp_pdf_judger" / "domain_details" / "overall.md").write_text(
        "# 전체 공통\n- 시험기준과 시험결과를 비교한다.\n",
        encoding="utf-8",
    )
    pdf_path = project_root / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nregression-test")
    summary_path = project_root / "summary_after.csv"
    summary_path.write_text(
        "제조명,검수합격,검수불합격,검수보류,총 계\n"
        "완제의약품,2,0,0,2\n"
        "전체,2,0,0,2\n",
        encoding="utf-8",
    )
    result = _make_result(pdf_path)
    for evaluation in result.evaluations:
        evaluation.final_status = PASS_LABEL
    result.summary = Summary(passed=3, failed=0, held=0, total=3)
    result.metadata.update(
        {
            "llm_enabled": True,
            "llm_call_count": 0,
            "llm_success_count": 0,
        }
    )
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )

    logged = append_rule_regression_log(
        result,
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        summary_counts_path=summary_path,
        rule_fingerprint_override="rules-v1",
    )

    workbook = load_workbook(logged.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
    }
    assert (
        sheet.cell(
            FIRST_RUN_ROW,
            headers["전체판정(충족/불충족/보류)"],
        ).value
        == "2/0/0"
    )
    assert sheet.cell(FIRST_RUN_ROW, headers["시험기준/결과 비교"]).value == "3/0/0"
    assert "LLM호출설명" not in headers
    assert "화면집계출처" not in headers
    assert "제조단계 미매핑 시험행" not in headers


def test_legacy_offset_summary_table_is_migrated_to_column_a(tmp_path, monkeypatch):
    project_root = tmp_path / "project"
    (project_root / "sp_pdf_judger" / "domain_details").mkdir(parents=True)
    (project_root / "sp_pdf_judger" / "domain_details" / "overall.md").write_text(
        "# 전체 공통\n- 시험기준과 시험결과를 비교한다.\n",
        encoding="utf-8",
    )
    pdf_path = project_root / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nregression-test")
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )

    first = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v1",
    )
    workbook = load_workbook(first.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    last_column = sheet.max_column
    last_row = sheet.max_row
    sheet.move_range(
        f"A4:{sheet.cell(last_row, last_column).coordinate}",
        cols=6,
        translate=False,
    )
    workbook.save(first.workbook_path)

    second = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v2",
    )
    workbook = load_workbook(second.workbook_path)
    sheet = workbook[SUMMARY_SHEET]

    assert sheet.cell(4, 1).value == "기록구분"
    assert sheet.cell(4, 7).value != "기록구분"
    assert sheet.cell(FIRST_RUN_ROW, 1).value == "실행결과"
    assert sheet.cell(FIRST_RUN_ROW + 1, 1).value == "실행결과"


def test_legacy_auto_draft_and_obsolete_columns_are_migrated(tmp_path, monkeypatch):
    project_root = tmp_path / "project"
    (project_root / "sp_pdf_judger" / "domain_details").mkdir(parents=True)
    (project_root / "sp_pdf_judger" / "domain_details" / "overall.md").write_text(
        "# 전체 공통\n- 시험기준과 시험결과를 비교한다.\n",
        encoding="utf-8",
    )
    pdf_path = project_root / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nregression-test")
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )
    first = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v1",
    )
    workbook = load_workbook(first.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
    }
    sheet.cell(GOLD_ROW, headers["정답지상태"], "초안(첫 실행값)")
    sheet.cell(GOLD_ROW, headers["시험기준/결과 비교"], "1/1/1")
    for header in (
        "PDF_SHA256",
        "실행_ID",
        "규칙코드지문",
        "RAG지문",
        "MD지문",
        "LLM호출",
        "LLM성공",
        "문서 구조/추적성",
        "참고 문서 적용",
        "보류 원칙",
        "LLM호출설명",
        "화면집계출처",
        "제조단계 미매핑 시험행",
    ):
        column = sheet.max_column + 1
        sheet.cell(4, column, header)
        sheet.cell(GOLD_ROW, column, "legacy")
        sheet.cell(FIRST_RUN_ROW, column, "legacy")
    workbook.save(first.workbook_path)

    second = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v2",
    )
    workbook = load_workbook(second.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
    }

    assert sheet.cell(GOLD_ROW, headers["정답지상태"]).value == "입력 대기"
    assert sheet.cell(GOLD_ROW, headers["시험기준/결과 비교"]).value is None
    assert "PDF_SHA256" not in headers
    assert "문서 구조/추적성" not in headers
    assert "LLM호출설명" not in headers
    assert "화면집계출처" not in headers
    assert "제조단계 미매핑 시험행" not in headers
    assert sheet.max_row == FIRST_RUN_ROW + 1


def test_new_rule_fingerprint_appends_and_manual_gold_is_preserved(tmp_path, monkeypatch):
    project_root = tmp_path / "project"
    (project_root / "sp_pdf_judger" / "domain_details").mkdir(parents=True)
    (project_root / "sp_pdf_judger" / "domain_details" / "overall.md").write_text(
        "# 전체 공통\n- 제조일자를 비교한다.\n",
        encoding="utf-8",
    )
    pdf_path = project_root / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nregression-test")
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )

    first = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v1",
    )
    workbook = load_workbook(first.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
    }
    sheet.cell(GOLD_ROW, headers["시험기준/결과 비교"], "99/0/0")
    workbook.save(first.workbook_path)

    second = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v2",
    )
    workbook = load_workbook(second.workbook_path)
    sheet = workbook[SUMMARY_SHEET]

    assert second.appended is True
    assert sheet.max_row == FIRST_RUN_ROW + 1
    assert sheet.cell(GOLD_ROW, headers["시험기준/결과 비교"]).value == "99/0/0"
    comparison_formula = sheet.cell(
        FIRST_RUN_ROW + 1,
        headers["정답지대비"],
    ).value
    assert comparison_formula.startswith("=")
    assert "시험기준/결과 비교" in comparison_formula


def test_llm_failure_and_success_are_separate_runs(tmp_path, monkeypatch):
    pdf_path = tmp_path / "sample.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n")
    project_root = tmp_path / "project"
    output_dir = project_root / "OUTPUT"
    overall_path = project_root / "sp_pdf_judger" / "domain_details" / "overall.md"
    overall_path.parent.mkdir(parents=True)
    overall_path.write_text("# 전체\n- 시험기준과 시험결과를 비교한다.\n", encoding="utf-8")
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )

    failed_result = _make_result(pdf_path)
    failed_result.metadata.update(
        {
            "llm_model": "gemini-test",
            "llm_call_count": 2,
            "llm_success_count": 0,
            "llm_last_error": "network unavailable",
        }
    )
    successful_result = _make_result(pdf_path)
    successful_result.metadata.update(
        {
            "llm_model": "gemini-test",
            "llm_call_count": 2,
            "llm_success_count": 2,
            "llm_last_error": "",
        }
    )

    failed = append_rule_regression_log(
        failed_result,
        project_root=project_root,
        output_dir=output_dir,
        rule_fingerprint_override="rules-a",
    )
    succeeded = append_rule_regression_log(
        successful_result,
        project_root=project_root,
        output_dir=output_dir,
        rule_fingerprint_override="rules-a",
    )

    assert failed.run_id != succeeded.run_id
    workbook = load_workbook(succeeded.workbook_path)
    assert workbook[SUMMARY_SHEET].max_row == FIRST_RUN_ROW + 1


def test_set_gold_from_latest_copies_only_result_columns(tmp_path, monkeypatch):
    project_root = tmp_path / "project"
    (project_root / "sp_pdf_judger" / "domain_details").mkdir(parents=True)
    (project_root / "sp_pdf_judger" / "domain_details" / "overall.md").write_text(
        "# 전체 공통\n- 시험기준과 시험결과를 비교한다.\n",
        encoding="utf-8",
    )
    pdf_path = project_root / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nregression-test")
    monkeypatch.setattr(
        "sp_pdf_judger.rule_regression._manufacturing_observations",
        _fake_manufacturing_observations,
    )
    logged = append_rule_regression_log(
        _make_result(pdf_path),
        project_root=project_root,
        output_dir=project_root / "OUTPUT",
        rule_fingerprint_override="rules-v1",
    )

    set_gold_from_latest(logged.workbook_path)
    summary = workbook_summary(logged.workbook_path)
    workbook = load_workbook(logged.workbook_path)
    sheet = workbook[SUMMARY_SHEET]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
    }

    assert summary["gold_status"] == "확정"
    assert sheet.cell(GOLD_ROW, headers["시험기준/결과 비교"]).value == "1/1/1"
    assert sheet.cell(GOLD_ROW, headers["제조번호 비교"]).value == "2/0/0"


def test_overall_catalog_exposes_unimplemented_page_rule(tmp_path):
    overall_md = tmp_path / "overall.md"
    overall_md.write_text(
        "# 전체 공통\n"
        "- 시험기준과 시험결과를 비교한다.\n"
        "- 페이지 번호가 순차적으로 이어지는지 검증한다.\n",
        encoding="utf-8",
    )

    entries = parse_overall_rule_catalog(overall_md)
    page_entry = next(entry for entry in entries if "페이지 번호" in entry.text)

    assert page_entry.executor == "page_number_validator"
    assert page_entry.implementation_status == "미구현"
